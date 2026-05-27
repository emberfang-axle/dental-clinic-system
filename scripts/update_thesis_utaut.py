"""Update UTAUT sections in the thesis docx (n=31, WM=4.24)."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "3k_DENTAL-CLINIC-APPPOINTMENT-AND-BILLING-SYSTEM.docx"

CONCLUSION_OLD = (
    "The UTAUT evaluation results (Overall WM = 4.52, Excellent) confirm that the system is well-accepted "
    "by its users, with strong ratings in Performance Expectancy and Behavioral Intention."
)
CONCLUSION_NEW = (
    "The UTAUT evaluation results (Overall WM = 4.24, Very Satisfactory, n = 31 respondents) confirm that "
    "the system is well-accepted by its users, with the highest ratings in Behavioral Intention (WM = 4.39) "
    "and Social Influence (WM = 4.27)."
)

APPENDIX_I_OLD = (
    "(The complete statistical analysis results including weighted means, standard deviations, "
    "and interpretation for all 50 respondents)"
)
APPENDIX_I_NEW = (
    "(Complete statistical analysis for n = 31 respondents: respondent profile, item-level weighted means, "
    "construct summaries, standard deviations, and interpretation — see tables in Section 10.3 and "
    "docs/UTAUT_RESULTS.md in the project repository.)"
)

NARRATIVE = (
    "The UTAUT evaluation was conducted with thirty-one (31) respondents, including two (2) staff members "
    "and twenty-nine (29) patients from Estandarte Dental Clinic and online respondents who tested the system. "
    "The overall weighted mean was 4.24, which is interpreted as Very Satisfactory. This indicates that users "
    "generally accept the system and find it useful for their daily tasks. The highest rated construct was "
    "Behavioral Intention (WM = 4.39), showing that users intend to continue using the system and would "
    "recommend it to others. Social Influence (WM = 4.27) indicates support from supervisors and colleagues. "
    "Performance Expectancy (WM = 4.19) and Effort Expectancy (WM = 4.16) confirm that users find the system "
    "useful and easy to use. Facilitating Conditions (WM = 4.21) shows that users have adequate resources and "
    "support to use the system effectively. Overall, the results confirm that the Estandarte Dental Clinic "
    "Appointment and Billing System meets user expectations."
)

TABLE_10_3 = [
    ["UTAUT Construct", "Weighted Mean", "Standard Deviation", "Descriptive Rating"],
    ["Performance Expectancy (PE)", "4.19", "0.31", "Very Satisfactory"],
    ["Effort Expectancy (EE)", "4.16", "0.30", "Very Satisfactory"],
    ["Social Influence (SI)", "4.27", "0.26", "Very Satisfactory"],
    ["Facilitating Conditions (FC)", "4.21", "0.29", "Very Satisfactory"],
    ["Behavioral Intention (BI)", "4.39", "0.32", "Very Satisfactory"],
    ["OVERALL WEIGHTED MEAN", "4.24", "0.30", "Very Satisfactory"],
]

PROFILE = [
    ["Category", "Count"],
    ["Role — Staff", "2"],
    ["Role — Patient", "29"],
    ["Age — 18-25", "27"],
    ["Age — 26-35", "2"],
    ["Age — 36-45", "2"],
    ["Gender — Male", "15"],
    ["Gender — Female", "16"],
]


def insert_paragraph_after(paragraph, text: str, bold: bool = False):
    new_p = deepcopy(paragraph._element)
    paragraph._element.addnext(new_p)
    from docx.text.paragraph import Paragraph

    p = Paragraph(new_p, paragraph._parent)
    run = p.add_run(text)
    if bold:
        run.bold = True
    run.font.size = Pt(11)
    return p


def add_table_after(paragraph, rows: list[list[str]]):
    parent = paragraph._parent
    # Create table via document body API
    doc = parent.part.document if hasattr(parent, "part") else paragraph.part.document
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
    paragraph._element.addnext(table._tbl)
    return table


def main():
    doc = Document(DOCX)

    # Update conclusion and appendix
    for p in doc.paragraphs:
        if CONCLUSION_OLD.split(".")[0] in p.text:
            p.text = p.text.replace(
                "The UTAUT evaluation results (Overall WM = 4.52, Excellent) confirm that the system is well-accepted by its users, with strong ratings in Performance Expectancy and Behavioral Intention.",
                CONCLUSION_NEW,
            )
        if "Overall WM = 4.52" in p.text:
            p.text = p.text.replace("Overall WM = 4.52, Excellent", "Overall WM = 4.24, Very Satisfactory (n = 31)")
        if APPENDIX_I_OLD in p.text:
            p.text = APPENDIX_I_NEW
        if "all 50 respondents" in p.text:
            p.text = p.text.replace("all 50 respondents", "all 31 respondents")

    # Insert UTAUT block after "Table 10.3 UTAUT Final Results Summary"
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip() == "Table 10.3 UTAUT Final Results Summary":
            anchor = p
            break

    if anchor is None:
        print("WARNING: Table 10.3 heading not found")
    else:
        # Skip if already updated
        nxt = anchor._element.getnext()
        if nxt is not None and "thirty-one" in "".join(
            t.text or "" for t in nxt.iter() if t.tag.endswith("}t")
        ):
            print("UTAUT block already present; skipping insert")
        else:
            p1 = insert_paragraph_after(anchor, "Respondent Profile Summary (n = 31)", bold=True)
            add_table_after(p1, PROFILE)
            p2 = insert_paragraph_after(p1, "")
            p3 = insert_paragraph_after(p2, "Table 10.3 presents the final UTAUT construct ratings.", bold=False)
            add_table_after(p3, TABLE_10_3)
            insert_paragraph_after(p3, NARRATIVE)

    doc.save(DOCX)
    print(f"Updated: {DOCX}")


if __name__ == "__main__":
    main()
