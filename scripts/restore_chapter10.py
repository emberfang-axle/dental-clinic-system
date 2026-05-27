"""Restore deleted Chapter 10.4–Appendices and insert UTAUT §10.3 content."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "3k_DENTAL-CLINIC-APPPOINTMENT-AND-BILLING-SYSTEM.docx"
OUT = ROOT / "3k_DENTAL-CLINIC-APPPOINTMENT-AND-BILLING-SYSTEM-UPDATED.docx"

NARRATIVE = (
    "The UTAUT evaluation was conducted with thirty-one (31) respondents, including two (2) staff members "
    "and twenty-nine (29) patients from Estandarte Dental Clinic and online respondents who tested the system. "
    "The overall weighted mean was 4.24, which is interpreted as Very Satisfactory. This indicates that users "
    "generally accept the system and find it useful for their daily tasks. The highest rated construct was "
    "Behavioral Intention (WM = 4.39), showing that users intend to continue using the system and would "
    "recommend it to others. Social Influence (WM = 4.27) indicates support from supervisors and colleagues. "
    "Performance Expectancy (WM = 4.19) and Effort Expectancy (WM = 4.16) confirm that users find the system "
    "useful and easy to use. Facilitating Conditions (WM = 4.21) shows that users have adequate resources and "
    "support to use the system effectively. Overall, the results confirm that the Estandarte Dental Clinic "
    "Appointment and Billing System meets user expectations."
)

TABLE_10_3 = [
    ["UTAUT Construct", "Weighted Mean", "Standard Deviation", "Descriptive Rating"],
    ["Performance Expectancy (PE)", "4.19", "0.31", "Very Satisfactory"],
    ["Effort Expectancy (EE)", "4.16", "0.30", "Very Satisfactory"],
    ["Social Influence (SI)", "4.27", "0.26", "Very Satisfactory"],
    ["Facilitating Conditions (FC)", "4.21", "0.29", "Very Satisfactory"],
    ["Behavioral Intention (BI)", "4.39", "0.32", "Very Satisfactory"],
    ["OVERALL WEIGHTED MEAN", "4.24", "0.30", "Very Satisfactory"],
]

PROFILE = [
    ["Category", "Count"],
    ["Role — Staff", "2"],
    ["Role — Patient", "29"],
    ["Age — 18-25", "27"],
    ["Age — 26-35", "2"],
    ["Age — 36-45", "2"],
    ["Gender — Male", "15"],
    ["Gender — Female", "16"],
]

RESTORE_PARAS = [
    "Table 10.3a — Respondent Profile Summary (n = 31)",
    "Table 10.3b — UTAUT Construct Ratings (n = 31)",
    NARRATIVE,
    "10.4 Recommendations",
    "Based on the findings of this study and the feedback received from UTAUT respondents, the following recommendations are proposed for future enhancements and improvements to the Estandarte Dental Clinic Appointment and Billing System.",
    "10.4.1 Technical Recommendations",
    "Mobile Application Development — Develop native iOS and Android applications using React Native to provide a better mobile experience for patients and staff, especially for those with limited internet connectivity.",
    "Offline Mode Enhancement — Expand the existing Firestore offline persistence to allow full offline functionality for staff, enabling queue management and patient lookup even without internet connection.",
    "Teleconsultation Module — Add video call functionality for remote consultations, allowing patients to consult with dentists without visiting the clinic physically.",
    "Electronic Medical Record (EMR) Integration — Explore integration with Philippine health information systems for standardized electronic medical records.",
    "10.4.2 Functional Recommendations",
    "Loyalty and Rewards Program — Implement a patient loyalty program with points accumulation, discounts for frequent visits, and referral rewards.",
    "Advanced Analytics Dashboard — Add predictive analytics for patient no-shows (using historical data), revenue forecasting, and inventory management for dental supplies.",
    "Multi-Branch Support — Extend the system to handle multiple clinic locations if Estandarte Dental Clinic expands to additional branches.",
    "Insurance Integration — Integrate with health insurance providers for direct billing and claims processing.",
    "Automated Appointment Reminder Customization — Allow patients to choose their preferred reminder time (e.g., 24 hours, 12 hours, 2 hours before appointment).",
    "10.4.3 Organizational Recommendations",
    "Regular Staff Training — Conduct quarterly refresher training sessions for staff to ensure they are utilizing all system features effectively.",
    "Data Backup Verification — Implement monthly verification of Firestore backups to ensure data can be restored if needed.",
    "User Feedback Collection — Continue collecting user feedback through the in-app feedback form and conduct annual UTAUT evaluations to measure ongoing user satisfaction.",
    "Security Audit — Conduct annual security audits and penetration testing to identify and address potential vulnerabilities.",
    "Documentation Update — Maintain and update system documentation including user manuals, API documentation, and deployment guides as the system evolves.",
    "10.5 Conclusion",
    "The Estandarte Dental Clinic Appointment and Billing System has successfully transformed the clinic's operations from a manual, paper-based process to a modern, digital, and automated platform. The system has effectively addressed the identified problems of lack of centralized records, scheduling delays, billing errors, and difficulty generating reports.",
    "The UTAUT evaluation results (Overall WM = 4.24, Very Satisfactory, n = 31 respondents) confirm that the system is well-accepted by its users, with the highest ratings in Behavioral Intention (WM = 4.39) and Social Influence (WM = 4.27). The system has improved appointment scheduling efficiency, automated billing processes, enhanced patient record management, and provided real-time reporting capabilities for clinic administration.",
    "The successful implementation of this system serves as a model for other small to medium-sized dental clinics in Compostela, Davao de Oro, and the broader Philippines region, demonstrating how modern web technologies (React, Firebase) can be leveraged to improve healthcare service delivery at an affordable cost.",
    "REFERENCES",
    "D. Bridge, \"Asia-Pacific Dental Practice Management Software Market Size, Share and Trends Analysis Report – Industry Overview and Forecast to 2032,\" (2023).",
    "E. Ramos, E. Bancud, S. Gumarang, M. Kummer, \"Dental Clinic Appointment Management System,\" (2025).",
    "R. Alejandro and F. Pajota, \"Implementation of online appointment and billing systems in private dental clinics,\" 2023.",
    "J. Manalo, Automation in healthcare management: Improving efficiency and patient satisfaction, 2024.",
    "K. Zoura, Automated billing system explained, 2023.",
    "A. Dhumal and P. Nikam, \"Paperless billing in healthcare: QR code and barcode systems,\" 2021.",
    "R. Pandey, V. Singh, and S. Kumar, \"Integrated digital billing and appointment systems in healthcare,\" 2024.",
    "S. Halde, P. Sharma, and A. Gupta, \"Efficiency improvements through integration of billing and scheduling,\" 2025.",
    "S. Baswaraju, R. Prasad, and K. Rajan, \"Implementation of improved billing system,\" 2020.",
    "L. Schwarz, \"Healthcare data security best practices,\" 2025.",
    "M. Laura and T. Manager, \"Cloud-based solutions for small healthcare clinics,\" 2024.",
    "F. Alshameri, R. Khan, and A. Yousuf, \"Integration of digital appointment, billing, and patient records in healthcare,\" 2025.",
    "X. Zhu, H. Li, and J. Wang, \"Integrated appointment and billing systems in dental clinics,\" 2024.",
    "International Journal of Electrical and Electronics Engineering, \"IoT in healthcare billing systems,\" 2024.",
    "R. Nithya, P. Kannan, and A. Singh, \"Digitalization of billing platforms in dental clinics,\" 2023.",
    "M. Taslim, A. Rahman, and S. Khan, \"Automation in dental clinic financial management,\" 2023.",
    "R. Anwat, S. Fatima, and H. Lee, \"Cloud-based appointment and billing systems in dental practice,\" 2024.",
    "D. Priyanshu, \"Improving data management in healthcare,\" 2025.",
    "APPENDICES",
    "Appendix A — Letter to Conduct the Study",
    "Appendix B — Letter to Conduct Interview",
    "Appendix C — Interview Guide Questions",
    "Appendix D — Patient Survey Questionnaire",
    "Appendix E — Staff/Owner Survey Questionnaire",
    "Appendix F — Logbook",
    "Appendix G — Documentation",
    "Appendix H — UTAUT Survey Questionnaire",
    "Appendix I — UTAUT Data Analysis Results",
    "(Complete statistical analysis for n = 31 respondents: respondent profile, item-level weighted means, construct summaries, standard deviations, and interpretation. See Section 10.3 and docs/UTAUT_RESULTS.md in the project repository.)",
    "Appendix J — Sample Screenshots of the System",
    "(Screenshots of the Admin Dashboard, Doctor Dashboard, Staff Dashboard, Patient Dashboard, Appointment Booking Wizard, Staff Appointment Queue, Clinical Records Interface, and Payment screens.)",
    "Appendix K — Deployment Guide",
    "(Step-by-step guide for deploying the system to Firebase Hosting, including environment variable configuration and Cloud Functions deployment.)",
    "Appendix L — User Manual",
    "(Comprehensive user manual for Admin, Doctor, Co-Doctor, Staff, and Patient roles, including how to perform common tasks in the system.)",
]


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
    return table


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    if bold:
        run.bold = True
    return p


def main():
    doc = Document(DOCX)

    # Only restore if Chapter 10 was truncated
    has_10_4 = any(p.text.strip().startswith("10.4 Recommendations") for p in doc.paragraphs)
    if has_10_4:
        print("Chapter 10.4 already present; skip restore")
        return

    add_para(doc, "Table 10.3a — Respondent Profile Summary (n = 31)", bold=True)
    add_table(doc, PROFILE)
    add_para(doc, "Table 10.3b — UTAUT Construct Ratings (n = 31)", bold=True)
    add_table(doc, TABLE_10_3)
    add_para(doc, NARRATIVE)

    for text in RESTORE_PARAS[3:]:  # skip duplicate headers already added
        bold = text.startswith("10.") or text == "REFERENCES" or text == "APPENDICES"
        add_para(doc, text, bold=bold)

    target = OUT
    try:
        doc.save(DOCX)
        target = DOCX
    except PermissionError:
        doc.save(OUT)
        print("NOTE: Original docx is locked (close Word). Saved to:", OUT)
    print("Restored Chapter 10.4–Appendices:", target)


if __name__ == "__main__":
    main()
